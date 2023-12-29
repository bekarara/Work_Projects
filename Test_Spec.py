from tkinter.filedialog import askopenfilename
import tkinter as tk
import xlsxwriter
import os
import pandas as pd
from docx import Document
from openpyxl import Workbook
from openpyxl import load_workbook
from pathlib import Path
from datetime import datetime
#========================================================#
def convert_doc_table_to_excel(word_doc_filepath, output_filepath):
    # read the test spec document
    doc = Document(word_doc_filepath)

    #convert to df
    tables = []
    for table in doc.tables:
        num_rows = len(table.rows)
        num_cols = len(table.columns)
        df = pd.DataFrame(index=range(num_rows), columns = range(num_cols))

        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                df.iloc[i,j] = cell.text

        #keep cells with unique values only:
        df = df.apply(lambda x: x.drop_duplicates(keep = 'first'),axis=1)

        #filter for test case tables only:
        if "NGR_TC_SI_LAB" in (table.rows[0].cells[0].text or table.rows[1].cells[0].text):
            df.iloc[12,4] = 'Test Result'
            df.iloc[12,5] = 'Remarks'
            for i in range(13,df.shape[0]): #df.shape[0] gives number of rows
                df.iloc[i,0] = i-12
            tables.append(pd.DataFrame(df))

    #write to excel
    writer = pd.ExcelWriter(output_filepath, engine = 'xlsxwriter')
    for df in tables:
        # write dataframes to excel
        df.to_excel(writer, sheet_name = df.iloc[0][0][18:22])
        
        #===========formatting===================================#
        workbook = writer.book
        #name sheets with test case ID
        worksheet = writer.sheets[df.loc[0][0][18:22]]
        #wrap text and set column width
        format_text_wrap = workbook.add_format({'text_wrap' : True})
        writer.sheets[df.loc[0][0][18:22]].set_column(1,6,40,format_text_wrap)
        # colour header row
        format_cell_colour = workbook.add_format({'bold' : True, 'bg_color' : 'silver'})
        #freeze pane
        worksheet.freeze_panes(14,0)
        #group the rows that are not test steps
        for freeze_row in range(2,13):
            worksheet.set_row(freeze_row, None, None, {'level' : 2, "hidden": True})
        worksheet.set_row(freeze_row+2, None, None,{'collapsed' : True})
        #colour header cells
        worksheet.set_row(13,None,format_cell_colour)
        
    writer.close()
    return
#========================================================#

# select input word doc file:
word_doc_filepath = askopenfilename()

#select directory for output:
root = tk.Tk()
root.withdraw()
output_dir = tk.filedialog.askdirectory()
filename_base = Path(word_doc_filepath).stem
#add timestamp and input file name to create output xlsx file name:
now = datetime.now()
timestamp = now.strftime('%Y-%m-%d-%H-%M-%S')
output_filepath = os.path.join(output_dir, filename_base + "_" + timestamp + ".xlsx")

# convert to excel:
convert_doc_table_to_excel(word_doc_filepath, output_filepath)